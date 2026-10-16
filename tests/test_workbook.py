"""Shared case workbook: template structure, import, and the ingest-page API.

The workbook is built by the code under test and filled in here with openpyxl,
the way a colleague would fill it in Excel. MRNs are made up.
"""

import io

import pytest
from dicom_series import write_phantom_series
from openpyxl import Workbook, load_workbook
from test_library import _scan, clinical, local  # noqa: F401  (local is a fixture)

from stoneplanner_api.library import validate_clinical
from stoneplanner_api.workbook import (
    COLUMNS,
    STATUS_HEADER,
    WorkbookError,
    build_template,
    parse_workbook,
)

COL = {c.key: i for i, c in enumerate(COLUMNS, start=1)}

COMPLETE = {
    "ct_site": "outside", "laterality": "Right", "or_position": "prone", "n_tracts": 2,
    "tract1_calyx": "lower", "tract1_approach": "subcostal",
    "tract2_calyx": "upper", "tract2_approach": "supracostal",
    "loc:renal pelvis": "yes", "loc:lower calyx": "yes",
    "staghorn": "no", "largest_mm": 28, "stone_free": "no", "residual_mm": 4,
    "postop_imaging": "KUB", "pleural_injury": "no", "colon_injury": "no",
    "transfusion": "yes", "age": 95, "sex": "female", "height_cm": 165, "weight_kg": 85,
}


def fill(key=None, rows=None, template=None) -> bytes:
    wb = load_workbook(io.BytesIO(template or build_template()))
    for r, (case_id, mrn) in enumerate((key or {}).items(), start=2):
        wb["Key"].cell(row=r, column=1, value=case_id)
        wb["Key"].cell(row=r, column=2, value=mrn)
    for r, values in enumerate(rows or [], start=3):
        for k, v in values.items():
            wb["Cases"].cell(row=r, column=COL[k], value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_template_structure():
    wb = load_workbook(io.BytesIO(build_template()))
    assert wb.sheetnames == ["Instructions", "Key", "Lookup", "Cases", "Lists"]
    assert wb["Lists"].sheet_state == "hidden"
    headers = [c.value for c in wb["Cases"][1]]
    assert headers == [c.header for c in COLUMNS] + [STATUS_HEADER]
    assert wb["Cases"].cell(row=2, column=1).value == "EXAMPLE"
    assert wb["Key"]["A2"].value == "case_001" and wb["Key"]["B2"].number_format == "@"
    assert "StudyIDs" in wb.defined_names
    # Every dropdown and number range on the Cases sheet has its validation.
    assert len(wb["Cases"].data_validations.dataValidation) == len(COLUMNS)
    # The Cases sheet itself never asks for an MRN.
    assert not any("mrn" in (h or "").lower() for h in headers)


def test_complete_row_imports_and_validates():
    data = fill({"case_001": "00123"}, [{"case_id": "case_001", **COMPLETE}])
    out = parse_workbook(data)
    assert out["problems"] == []
    (row,) = out["rows"]                          # the example row is skipped
    assert row["case_id"] == "case_001" and row["mrn"] == "00123"
    assert row["errors"] == [] and row["postop_imaging"] == "kub"
    c = row["clinical"]
    assert c["ct"]["position"] == "supine"        # blank means supine
    assert c["operative"]["laterality"] == "right"
    assert c["stone"]["staghorn"] == "none"
    assert c["stone"]["location"] == ["lower calyx", "renal pelvis"]
    assert c["demographics"]["bmi"] == 31.2       # from height and weight
    rec = validate_clinical(c)
    assert rec["demographics"]["age"] == 90 and rec["postop"]["residual_mm"] == 4.0


def test_incomplete_rows_say_what_is_missing():
    rows = [
        {"case_id": "case_001", **COMPLETE, "tract2_approach": None, "residual_mm": None,
                                   "bmi": None, "height_cm": None},
        {"case_id": "case_002", **COMPLETE},
        {"case_id": "case_001", **COMPLETE},
        {"case_id": "case_003", **COMPLETE, "largest_mm": "big"},
    ]
    out = parse_workbook(fill({"case_001": 4567, "case_003": "x9"}, rows))
    first, second, dup, bad = out["rows"]
    assert first["mrn"] == "4567"                 # a number in the Key comes back as text
    assert {"tract 2", "residual fragment size", "BMI (or height and weight)"} <= set(
        first["errors"])
    assert second["errors"][0] == "no MRN for this study ID on the Key sheet"
    assert dup["errors"][0] == "study ID also used on row 3"
    assert any("number" in e for e in bad["errors"])


def test_duplicate_mrn_in_key_is_reported():
    out = parse_workbook(fill({"case_001": "111", "case_002": "111"}))
    assert "case_001, case_002" in out["problems"][0]


@pytest.mark.parametrize("data, message", [
    (b"not a spreadsheet", "could not read"),
    (None, "no Key sheet"),
])
def test_not_a_case_workbook(data, message):
    if data is None:
        buf = io.BytesIO()
        Workbook().save(buf)
        data = buf.getvalue()
    with pytest.raises(WorkbookError, match=message):
        parse_workbook(data)


# ---------------------------------------------------------------------------
# Through the API


def _import(client, data: bytes) -> dict:
    r = client.post("/api/ingest/workbook", files={"file": ("book.xlsx", data)})
    assert r.status_code == 200, r.text
    return r.json()


def test_template_download_and_import_statuses(local):  # noqa: F811
    client, tmp = local
    r = client.get("/api/ingest/workbook.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert "stone-planner-cases.xlsx" in r.headers["content-disposition"]

    book = fill({"case_001": "MRN1", "case_002": "MRN2"},
                [{"case_id": "case_001", **COMPLETE},
                 {"case_id": "case_002", **COMPLETE, "sex": None}],
                template=r.content)
    rows = _import(client, book)["rows"]
    assert [r["status"] for r in rows] == ["ready", "incomplete"]

    # Ingest row 1 the way the page does: its clinical data plus a scanned CT.
    write_phantom_series(tmp / "study")
    pre = _scan(client, tmp / "study")
    first = rows[0]
    r = client.post("/api/ingest/cases", json={
        "case_id": first["case_id"], "mrn": first["mrn"], "clinical": first["clinical"],
        "preop": {"staging": pre["staging"], "series_uid": pre["recommended_ct"]}})
    assert r.status_code == 200, r.text
    assert _import(client, book)["rows"][0]["status"] == "ingested"

    # The same study ID linked to another patient in the workbook is refused.
    other = fill({"case_001": "SOMEONE-ELSE"}, [{"case_id": "case_001", **COMPLETE}])
    row = _import(client, other)["rows"][0]
    assert row["status"] == "conflict" and "different MRN" in row["errors"][0]


def test_ingest_refuses_a_study_id_linked_to_another_mrn(local):  # noqa: F811
    client, tmp = local
    write_phantom_series(tmp / "study")
    pre = _scan(client, tmp / "study")
    body = {"case_id": "case_007", "mrn": "A1", "clinical": clinical(),
            "preop": {"staging": pre["staging"], "series_uid": pre["recommended_ct"]}}
    assert client.post("/api/ingest/cases", json=body).status_code == 200
    # Simulate the case folder being removed by hand while the key keeps the link.
    import shutil
    shutil.rmtree(tmp / "data" / "case_007")
    pre = _scan(client, tmp / "study")
    body.update(mrn="B2", preop={"staging": pre["staging"], "series_uid": pre["recommended_ct"]})
    r = client.post("/api/ingest/cases", json=body)
    assert r.status_code == 400 and "different MRN" in r.json()["detail"]


def test_workbook_routes_need_the_local_library(tmp_path):
    from fastapi.testclient import TestClient

    from stoneplanner_api.app import create_app
    from stoneplanner_api.config import Settings

    settings = Settings(secret_key="k" * 32, app_password="pw", cookie_secure=False,
                        work_dir=tmp_path / "work")
    with TestClient(create_app(settings)) as c:
        c.post("/api/login", json={"password": "pw"})
        assert c.get("/api/ingest/workbook.xlsx").status_code == 404
        assert c.post("/api/ingest/workbook",
                      files={"file": ("b.xlsx", b"x")}).status_code == 404
