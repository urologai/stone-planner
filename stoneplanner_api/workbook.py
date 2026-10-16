"""Shared case workbook: an Excel template for the clinical data, and its import.

The workbook is for the team's approved university server, where colleagues
fill it in together. It holds the MRN, so it never goes anywhere else. Sheets:

    Instructions  how to fill it in, and what every column means
    Key           Study ID and MRN, one row per patient: the MRN-to-study-ID key
    Lookup        type an MRN to find its study ID, or a study ID to find its MRN
    Cases         one row per case, keyed by Study ID, with no MRN on it
    Lists         dropdown values (hidden)

build_template() writes the empty workbook. parse_workbook() reads a filled
one back into rows shaped like the ingest form (POST /api/ingest/cases), each
with the problems found, so the ingest page can fill its form from a row. The
DICOM folders are still picked on the page; the workbook carries no images.

COLUMNS is the single description of the Cases sheet: the dropdowns, the
"Still needed" formula, the instructions table and the import all come from
it, and its choices come from library.py so the workbook and the form agree.

    python -m stoneplanner_api.workbook stone-planner-cases.xlsx
"""

from __future__ import annotations

import io
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from stoneplanner_api.library import GROUPS, STONE_LOCATIONS, LibraryError, validate_clinical

N_KEY_ROWS = 500           # study IDs pre-filled on the Key sheet
N_CASE_ROWS = 500          # Cases rows with dropdowns and the "Still needed" formula
EXAMPLE_ID = "EXAMPLE"
MAX_TRACTS = 6
FONT = "Arial"

YES_NO = ("yes", "no")


@dataclass(frozen=True)
class Col:
    key: str
    header: str
    kind: str                       # "id", "choice", "flag", "whole", "decimal"
    choices: tuple = ()
    lo: float | None = None
    hi: float | None = None
    required: str = "yes"           # "yes", "no", or a condition shown to the user
    where: str = ""
    example: object = None
    width: int = 14


def _columns() -> list[Col]:
    cols = [
        Col("case_id", "Study ID", "id", where="Key sheet: the study ID next to the MRN",
            example=EXAMPLE_ID, width=12),
        Col("ct_position", "CT patient position", "choice", ("supine", "prone", "other"),
            required="no (blank means supine)", where="CT report or scout", example="supine"),
        Col("ct_site", "CT done here or outside", "choice", ("here", "outside"),
            where="CT report header", example="outside"),
        Col("laterality", "Laterality", "choice", ("left", "right"), where="operative report",
            example="right", width=11),
        Col("or_position", "OR position", "choice", ("prone", "supine"),
            where="operative report", example="prone", width=11),
        Col("n_tracts", "Number of tracts", "whole", lo=1, hi=MAX_TRACTS,
            where="operative report", example=2, width=10),
    ]
    examples = {1: ("lower", "subcostal"), 2: ("upper", "supracostal")}
    for i in range(1, MAX_TRACTS + 1):
        calyx, approach = examples.get(i, (None, None))
        cols += [
            Col(f"tract{i}_calyx", f"Tract {i} calyceal group", "choice", tuple(GROUPS),
                required=f"if {i} or more tracts", where="operative report", example=calyx),
            Col(f"tract{i}_approach", f"Tract {i} approach", "choice",
                ("subcostal", "supracostal"), required=f"if {i} or more tracts",
                where="operative report", example=approach, width=12),
        ]
    cols += [
        Col("sheath_fr", "Sheath size (Fr)", "decimal", lo=6, hi=40, required="no",
            where="operative report", example=30, width=10),
        Col("guidance", "Access guidance", "choice", ("fluoroscopy", "ultrasound", "combined"),
            required="no", where="operative report", example="fluoroscopy", width=13),
    ]
    for loc in STONE_LOCATIONS:
        cols.append(Col(f"loc:{loc}", f"Stone: {loc}", "choice", ("yes",),
                        required="at least one location", where="CT read or op report",
                        example="yes" if loc in ("renal pelvis", "lower calyx") else None,
                        width=10))
    cols += [
        Col("staghorn", "Staghorn", "choice", ("no", "partial", "complete"),
            where="CT read or op report", example="partial", width=10),
        Col("largest_mm", "Largest stone dimension (mm)", "decimal", lo=1, hi=200,
            where="CT read", example=28, width=12),
        Col("manual_volume_mm3", "Manual stone volume (mm3)", "decimal", lo=1, hi=1e6,
            required="no", where="measured by the team", width=12),
        Col("stone_free", "Stone free", "choice", YES_NO, where="postop imaging read",
            example="no", width=9),
        Col("residual_mm", "Largest residual fragment (mm)", "decimal", lo=0, hi=200,
            required="if not stone free", where="postop imaging read", example=4, width=12),
        Col("postop_imaging", "Postop imaging to export", "choice", ("CT", "KUB", "none"),
            required="no", where="PACS; a reminder of which folder to pick", example="KUB",
            width=11),
        Col("pleural_injury", "Pleural injury", "choice", YES_NO, where="chart",
            example="no", width=9),
        Col("colon_injury", "Colon injury", "choice", YES_NO, where="chart",
            example="no", width=9),
        Col("transfusion", "Transfusion", "choice", YES_NO, where="chart", example="no",
            width=9),
        Col("age", "Age (years)", "whole", lo=0, hi=120, where="chart; 90 and over is kept "
            "as 90", example=61, width=8),
        Col("sex", "Sex", "choice", ("female", "male", "other"), where="chart",
            example="female", width=9),
        Col("bmi", "BMI", "decimal", lo=10, hi=90, required="or height and weight",
            where="chart", example=31.2, width=7),
        Col("height_cm", "Height (cm)", "decimal", lo=50, hi=250, required="no",
            where="chart, if BMI is not recorded", width=8),
        Col("weight_kg", "Weight (kg)", "decimal", lo=20, hi=400, required="no",
            where="chart, if BMI is not recorded", width=8),
    ]
    return cols


COLUMNS = _columns()
STATUS_HEADER = "Still needed"


# ---------------------------------------------------------------------------
# Template


THIN = Side(style="thin", color="C9C7BC")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="232F38")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")       # cells to type in
FORMULA_FILL = PatternFill("solid", fgColor="ECEBE6")     # filled by formula
EXAMPLE_FONT = Font(name=FONT, italic=True, color="7A858C")


def _font(**kw) -> Font:
    return Font(name=FONT, **kw)


def _header(ws, row: int, values: list[str], widths: list[int] | None = None) -> None:
    for i, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=value)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 44


def _list_name(col: Col) -> str:
    return "list_" + "".join(c if c.isalnum() else "_" for c in col.key)


def _still_needed_formula(row: int) -> str:
    """Excel formula naming the required fields still blank on one Cases row."""
    letter = {c.key: get_column_letter(i) for i, c in enumerate(COLUMNS, start=1)}

    def ref(key: str) -> str:
        return f"{letter[key]}{row}"

    parts = []
    for col in COLUMNS:
        if col.required == "yes" and col.kind != "id":
            parts.append(f'IF({ref(col.key)}="","{col.header}, ","")')
    for i in range(1, MAX_TRACTS + 1):
        calyx, approach = ref(f"tract{i}_calyx"), ref(f"tract{i}_approach")
        parts.append(f'IF(AND(N({ref("n_tracts")})>={i},OR({calyx}="",{approach}="")),'
                     f'"tract {i}, ","")')
    locs = [ref(f"loc:{loc}") for loc in STONE_LOCATIONS]
    parts.append(f'IF(COUNTIF({locs[0]}:{locs[-1]},"yes")=0,"stone location, ","")')
    parts.append(f'IF(AND({ref("stone_free")}="no",{ref("residual_mm")}=""),'
                 f'"residual fragment size, ","")')
    parts.append(f'IF(AND({ref("bmi")}="",OR({ref("height_cm")}="",{ref("weight_kg")}="")),'
                 f'"BMI, ","")')
    parts.append(f'IF(COUNTIFS(Key!$A:$A,{ref("case_id")},Key!$B:$B,"<>")=0,'
                 f'"an MRN for this study ID on the Key, ","")')
    expr = "&".join(parts)
    return (f'=IF({ref("case_id")}="","",IF({expr}="","complete",'
            f'LEFT({expr},LEN({expr})-2)))')


def build_template() -> bytes:
    wb = Workbook()
    _instructions(wb.active)
    key = wb.create_sheet("Key")
    lookup = wb.create_sheet("Lookup")
    cases = wb.create_sheet("Cases")
    lists = wb.create_sheet("Lists")
    lists.sheet_state = "hidden"

    # Lists: one column per dropdown, each a named range.
    choice_cols = [c for c in COLUMNS if c.kind == "choice"]
    for i, col in enumerate(choice_cols, start=1):
        letter = get_column_letter(i)
        lists.cell(row=1, column=i, value=col.header).font = _font(bold=True)
        for j, choice in enumerate(col.choices, start=2):
            lists.cell(row=j, column=i, value=choice).font = _font()
        ref = f"{quote_sheetname('Lists')}!${letter}$2:${letter}${len(col.choices) + 1}"
        wb.defined_names[_list_name(col)] = DefinedName(_list_name(col), attr_text=ref)

    _key_sheet(wb, key)
    _lookup_sheet(lookup)
    _cases_sheet(cases)
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _key_sheet(wb, ws) -> None:
    _header(ws, 1, ["Study ID", "MRN", "Added on", "On the Cases sheet"], [12, 18, 12, 18])
    ws.freeze_panes = "A2"
    for r in range(2, N_KEY_ROWS + 2):
        ws.cell(row=r, column=1, value=f"case_{r - 1:03d}").font = _font()
        mrn = ws.cell(row=r, column=2)
        mrn.number_format = "@"              # text, so leading zeros survive
        mrn.fill = INPUT_FILL
        mrn.font = _font()
        added = ws.cell(row=r, column=3)
        added.number_format = "yyyy-mm-dd"
        added.fill = INPUT_FILL
        added.font = _font()
        status = ws.cell(row=r, column=4,
                         value=f'=IF(B{r}="","",IF(COUNTIF(Cases!$A:$A,A{r})>0,"yes","not yet"))')
        status.fill = FORMULA_FILL
        status.font = _font(color="5D6B73")
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = BOX
    last = N_KEY_ROWS + 1
    wb.defined_names["StudyIDs"] = DefinedName("StudyIDs", attr_text=f"Key!$A$2:$A${last}")
    dv = DataValidation(type="custom", formula1='AND(LEN(B2)<=64,ISERROR(FIND(",",B2)))',
                        allow_blank=True, showErrorMessage=True, errorTitle="MRN",
                        error="Up to 64 characters, no commas.")
    dv.add(f"B2:B{last}")
    ws.add_data_validation(dv)
    dup = FormulaRule(formula=[f'AND(B2<>"",COUNTIF($B$2:$B${last},B2)>1)'],
                      fill=PatternFill("solid", fgColor="F6DFE2"),
                      font=Font(name=FONT, color="A33A4A", bold=True))
    ws.conditional_formatting.add(f"B2:B{last}", dup)
    ws["F1"] = "Type each patient's MRN next to the next free study ID. Never change a study " \
               "ID once it has an MRN. A duplicate MRN turns red."
    ws["F1"].font = _font(italic=True, color="5D6B73")


def _lookup_sheet(ws) -> None:
    last = N_KEY_ROWS + 1
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 24
    ws["A1"] = "MRN and study ID converter"
    ws["A1"].font = _font(bold=True, size=13)
    rows = [
        (3, "Type an MRN", None),
        (4, "Study ID",
         ('=IF(B3="","",IFERROR(INDEX(Key!$A$2:$A${l},MATCH(B3&"",Key!$B$2:$B${l},0)),'
         'IFERROR(INDEX(Key!$A$2:$A${l},MATCH(--B3,Key!$B$2:$B${l},0)),"not in the Key")))')),
        (6, "Type a study ID", None),
        (7, "MRN",
         ('=IF(B6="","",IFERROR(IF(INDEX(Key!$B$2:$B${l},MATCH(B6,Key!$A$2:$A${l},0))="",'
         '"no MRN yet",INDEX(Key!$B$2:$B${l},MATCH(B6,Key!$A$2:$A${l},0))&""),'
         '"not a study ID"))')),
    ]
    for row, label, formula in rows:
        ws.cell(row=row, column=1, value=label).font = _font(bold=formula is None)
        cell = ws.cell(row=row, column=2)
        cell.border = BOX
        cell.font = _font()
        if formula:
            cell.value = formula.format(l=last)
            cell.fill = FORMULA_FILL
        else:
            cell.number_format = "@"
            cell.fill = INPUT_FILL
    ws["A9"] = "Yellow cells are for typing; grey cells fill in by themselves."
    ws["A9"].font = _font(italic=True, color="5D6B73")


def _cases_sheet(ws) -> None:
    headers = [c.header for c in COLUMNS] + [STATUS_HEADER]
    widths = [c.width for c in COLUMNS] + [40]
    _header(ws, 1, headers, widths)
    ws.freeze_panes = "B2"
    first, last = 3, N_CASE_ROWS + 2

    for i, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        ex = ws.cell(row=2, column=i, value=col.example)
        ex.font = EXAMPLE_FONT
        ex.border = BOX
        rng = f"{letter}{first}:{letter}{last}"
        if col.kind == "id":
            dv = DataValidation(type="list", formula1="=StudyIDs", allow_blank=True)
        elif col.kind == "choice":
            dv = DataValidation(type="list", formula1=f"={_list_name(col)}", allow_blank=True)
        else:
            dv = DataValidation(type="whole" if col.kind == "whole" else "decimal",
                                operator="between", formula1=str(col.lo), formula2=str(col.hi),
                                allow_blank=True)
        dv.showErrorMessage = True
        dv.errorTitle = col.header
        dv.error = ("Choose from the list." if col.kind in ("id", "choice")
                    else f"Enter a number from {col.lo:g} to {col.hi:g}.")
        dv.add(rng)
        ws.add_data_validation(dv)
        for r in range(first, last + 1):
            cell = ws.cell(row=r, column=i)
            cell.fill = INPUT_FILL
            cell.font = _font()
            cell.border = BOX

    status_col = len(COLUMNS) + 1
    ws.cell(row=2, column=status_col, value="example row: the import skips it").font = \
        EXAMPLE_FONT
    for r in range(first, last + 1):
        cell = ws.cell(row=r, column=status_col, value=_still_needed_formula(r))
        cell.fill = FORMULA_FILL
        cell.font = _font(color="5D6B73")
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    s = get_column_letter(status_col)
    ws.conditional_formatting.add(
        f"{s}{first}:{s}{last}",
        FormulaRule(formula=[f'{s}{first}="complete"'], fill=PatternFill("solid", fgColor="DDEEEA"),
                    font=Font(name=FONT, color="0E7A72", bold=True)))
    ws.conditional_formatting.add(
        f"A{first}:A{last}",
        FormulaRule(formula=[f'AND(A{first}<>"",COUNTIF($A${first}:$A${last},A{first})>1)'],
                    fill=PatternFill("solid", fgColor="F6DFE2")))
    # Tracts 4 to 6 are rare: grouped so they can be folded away.
    first_extra = next(i for i, c in enumerate(COLUMNS, start=1) if c.key == "tract4_calyx")
    ws.column_dimensions.group(get_column_letter(first_extra),
                               get_column_letter(first_extra + 5), hidden=True)


def _instructions(ws) -> None:
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 40
    lines = [
        ("Stone planner case workbook", {"bold": True, "size": 14}),
        (("This workbook holds MRNs. Keep it on the approved university server. Do not "
         "email it, copy it to a laptop, or put it on a personal cloud drive."), {"bold": True}),
        ("", {}),
        (("1. Key sheet: type the patient's MRN next to the next free study ID. The study ID "
         "is the only identifier the planner uses. Never change a study ID once it has an MRN."),
         {}),
        ("2. Lookup sheet: type an MRN to find its study ID, or a study ID to find its MRN.", {}),
        (("3. Cases sheet: one row per case. Choose the study ID from the list, then fill the "
         "yellow cells from the dropdowns. Row 2 is an example and is skipped on import."), {}),
        (("4. The last column, Still needed, lists blank required fields until it says "
         "complete."), {}),
        (("5. To ingest: on the computer that runs the planner, open the Ingest page, press "
         "Import workbook and choose this file. Pick a row to fill the form, then choose the "
         "CT folder (and the postop folder, if any) and press Ingest case."), {}),
        ("", {}),
        (("Yellow cells are for typing. Grey cells fill in by themselves; leave them alone. "
         "The Cases sheet has no MRN, so a copy of it without the Key sheet is de-identified."),
         {"italic": True, "color": "5D6B73"}),
        ("", {}),
    ]
    for r, (text, style) in enumerate(lines, start=1):
        cell = ws.cell(row=r, column=1, value=text)
        cell.font = _font(**style)
        if text:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 30 if len(text) > 90 else 18
    start = len(lines) + 1
    _header(ws, start, ["Column (Cases sheet)", "Required", "Values", "Where to find it"])
    for r, col in enumerate(COLUMNS, start=start + 1):
        if col.kind in ("choice", "id"):
            values = ", ".join(col.choices) if col.choices else "from the Key sheet"
        else:
            values = f"{'whole number' if col.kind == 'whole' else 'number'}, " \
                     f"{col.lo:g} to {col.hi:g}"
        for c, value in enumerate((col.header, col.required, values, col.where), start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = _font()
            cell.border = BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")


# ---------------------------------------------------------------------------
# Import


class WorkbookError(LibraryError):
    """The file is not a usable case workbook."""


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _choice(value) -> str | None:
    text = _text(value).lower()
    return text or None


def _flag(value) -> bool | None:
    if isinstance(value, bool):
        return value
    text = _text(value).lower()
    if text in ("yes", "y", "true", "1"):
        return True
    if text in ("no", "n", "false", "0"):
        return False
    return None


def _num(value):
    if value is None or _text(value) == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return _text(value)          # left for validate_clinical to report


def _header_map(ws, wanted: list[str]) -> dict[str, int]:
    found = {}
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        name = _text(cell.value).lower()
        if name:
            found[name] = cell.column - 1
    missing = [h for h in wanted if h.lower() not in found]
    if missing:
        raise WorkbookError(f"{ws.title} sheet is missing columns: {', '.join(missing)}. Start "
                            "from the workbook the Ingest page provides.")
    return {h: found[h.lower()] for h in wanted}


def read_key(wb) -> tuple[dict[str, str], list[str]]:
    """Study ID to MRN from the Key sheet, and problems with it."""
    if "Key" not in wb.sheetnames:
        raise WorkbookError("no Key sheet. Start from the workbook the Ingest page provides.")
    ws = wb["Key"]
    cols = _header_map(ws, ["Study ID", "MRN"])
    key, problems, by_mrn = {}, [], {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        case_id = _text(row[cols["Study ID"]])
        mrn = _text(row[cols["MRN"]])
        if not case_id or not mrn:
            continue
        if case_id in key:
            problems.append(f"Key: study ID {case_id} appears twice")
            continue
        key[case_id] = mrn
        by_mrn.setdefault(mrn, []).append(case_id)
    for mrn, ids in by_mrn.items():
        if len(ids) > 1:
            problems.append(f"Key: one MRN is entered for {', '.join(ids)}; if that is the same "
                            "patient on purpose (other side, repeat PCNL), ignore this")
    return key, problems


def _row_clinical(get) -> tuple[dict, list[str]]:
    missing = []

    def need(label: str, value):
        if value in (None, ""):
            missing.append(label)
        return value

    n = _num(get("n_tracts"))
    n_tracts = int(n) if isinstance(n, float) and n.is_integer() and 1 <= n <= MAX_TRACTS \
        else None
    if n_tracts is None:
        missing.append("number of tracts (1 to 6)")
    tracts = []
    for i in range(1, (n_tracts or 0) + 1):
        calyx, approach = _choice(get(f"tract{i}_calyx")), _choice(get(f"tract{i}_approach"))
        if not calyx or not approach:
            missing.append(f"tract {i}")
        tracts.append({"calyx": calyx or "", "approach": approach or ""})

    locations = [loc for loc in STONE_LOCATIONS if _flag(get(f"loc:{loc}"))]
    if not locations:
        missing.append("stone location")
    staghorn = _choice(get("staghorn"))
    stone_free = _flag(get("stone_free"))
    residual = _num(get("residual_mm"))
    if stone_free is False and residual is None:
        missing.append("residual fragment size")
    bmi = _num(get("bmi"))
    height, weight = _num(get("height_cm")), _num(get("weight_kg"))
    if bmi is None and isinstance(height, float) and isinstance(weight, float) and height > 0:
        bmi = round(weight / (height / 100) ** 2, 1)

    clinical = {
        "ct": {"position": _choice(get("ct_position")) or "supine",
               "site": need("CT done here or outside", _choice(get("ct_site")))},
        "operative": {
            "laterality": need("laterality", _choice(get("laterality"))),
            "or_position": need("OR position", _choice(get("or_position"))),
            "tracts": tracts,
            "sheath_fr": _num(get("sheath_fr")),
            "guidance": _choice(get("guidance")),
        },
        "stone": {
            "location": locations,
            "staghorn": {"no": "none"}.get(staghorn, staghorn),
            "largest_mm": need("largest stone dimension", _num(get("largest_mm"))),
            "manual_volume_mm3": _num(get("manual_volume_mm3")),
        },
        "postop": {"stone_free": stone_free, "residual_mm": None if stone_free else residual},
        "complications": {k: _flag(get(k)) for k in
                          ("pleural_injury", "colon_injury", "transfusion")},
        "demographics": {"age": need("age", _num(get("age"))),
                         "sex": need("sex", _choice(get("sex"))),
                         "bmi": bmi},
    }
    if staghorn is None:
        missing.append("staghorn")
    if stone_free is None:
        missing.append("stone free")
    for k, label in (("pleural_injury", "pleural injury"), ("colon_injury", "colon injury"),
                     ("transfusion", "transfusion")):
        if clinical["complications"][k] is None:
            missing.append(label)
    if bmi is None:
        missing.append("BMI (or height and weight)")
    return clinical, missing


def parse_workbook(data: bytes) -> dict:
    """Rows of a filled workbook, each shaped like the ingest form, with its problems."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise WorkbookError("could not read that file as an Excel workbook (.xlsx)") from exc
    try:
        key, problems = read_key(wb)
        if "Cases" not in wb.sheetnames:
            raise WorkbookError("no Cases sheet. Start from the workbook the Ingest page "
                                "provides.")
        ws = wb["Cases"]
        cols = _header_map(ws, [c.header for c in COLUMNS])
        by_key = {c.key: cols[c.header] for c in COLUMNS}
        rows, seen = [], {}
        for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            case_id = _text(values[by_key["case_id"]])
            if not case_id or case_id.upper() == EXAMPLE_ID:
                continue

            def get(k, values=values):
                return values[by_key[k]]

            clinical, errors = _row_clinical(get)
            if case_id in seen:
                errors.insert(0, f"study ID also used on row {seen[case_id]}")
            seen.setdefault(case_id, excel_row)
            mrn = key.get(case_id)
            if not mrn:
                errors.insert(0, "no MRN for this study ID on the Key sheet")
            if not errors:
                try:
                    validate_clinical(clinical)
                except LibraryError as exc:
                    errors.append(str(exc))
            rows.append({"row": excel_row, "case_id": case_id, "mrn": mrn or "",
                         "clinical": clinical, "errors": errors,
                         "postop_imaging": _choice(get("postop_imaging"))})
    finally:
        wb.close()
    return {"rows": rows, "problems": problems}


def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    if len(args) != 1:
        print("usage: python -m stoneplanner_api.workbook OUT.xlsx", file=sys.stderr)
        return 2
    Path(args[0]).write_bytes(build_template())
    print(f"wrote {args[0]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
